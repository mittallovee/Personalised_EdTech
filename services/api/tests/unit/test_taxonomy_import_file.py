"""M3 PR 2c-1: reading an uploaded syllabus file safely (C-03, T25). Synthetic files only."""

import io
import re
import time
import tracemalloc
import zipfile

import pytest
from openpyxl import Workbook

from qbank.modules.taxonomy import import_file
from qbank.modules.taxonomy.import_file import ImportFileInvalid, read_table


def xlsx(*rows: list[object], extra_sheet: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    if extra_sheet:
        workbook.create_sheet("Other").append(["ignored"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def refused(data: bytes) -> str:
    with pytest.raises(ImportFileInvalid) as caught:
        read_table(data)
    assert caught.value.code == "IMPORT_INVALID"
    assert caught.value.status_code == 422
    return str(caught.value)


def test_a_spreadsheet_is_read_into_columns_and_rows() -> None:
    table = read_table(
        xlsx(
            ["Kind", "Path", "Sort"],
            ["CLASS", "Class 11", 1],
            ["SUBJECT", "Class 11 > Physics", 2.0],
        )
    )
    assert table.columns == ["kind", "path", "sort"]
    assert [(r.number, r.cells) for r in table.rows] == [
        (2, {"kind": "CLASS", "path": "Class 11", "sort": "1"}),
        (3, {"kind": "SUBJECT", "path": "Class 11 > Physics", "sort": "2"}),
    ]


def test_only_the_first_sheet_is_read() -> None:
    table = read_table(xlsx(["kind"], ["CLASS"], extra_sheet=True))
    assert [r.cells for r in table.rows] == [{"kind": "CLASS"}]


def test_a_csv_file_is_read_the_same_way_including_a_byte_order_mark() -> None:
    data = "﻿Kind,Path,label_hi\r\nCLASS,Class 11,कक्षा 11\r\n\r\nSUBJECT,Class 11 > Physics,\r\n"
    table = read_table(data.encode())
    assert table.columns == ["kind", "path", "label_hi"]
    # The blank line is skipped, but the row after it keeps the number the admin sees in the file.
    assert [(r.number, r.cells["path"], r.cells["label_hi"]) for r in table.rows] == [
        (2, "Class 11", "कक्षा 11"),
        (4, "Class 11 > Physics", ""),
    ]


def test_cells_are_trimmed_and_short_rows_are_filled_with_blanks() -> None:
    table = read_table(b"kind,path,sort\n  CLASS  ,Class 11\n")
    assert table.rows[0].cells == {"kind": "CLASS", "path": "Class 11", "sort": ""}


def test_a_spreadsheet_with_a_formula_is_refused() -> None:
    message = refused(xlsx(["kind", "path"], ["CLASS", '=CONCATENATE("Class", " 11")']))
    assert "formulas" in message


def test_the_file_is_judged_by_its_content_not_its_name() -> None:
    assert "UTF-8" in refused("kind,path\nCLASS,Cl\xe4ss".encode("latin-1"))
    assert "not a spreadsheet" in refused(b"\x89PNG\r\n\x1a\n\x00\x00\x00")
    assert "valid .xlsx" in refused(b"PK\x03\x04 this is not really an archive")
    assert ".xlsx" in refused(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 100)  # old Excel
    refused(b"")


def test_a_zip_that_is_not_a_spreadsheet_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hello")
    assert "valid .xlsx" in refused(buffer.getvalue())


def test_an_archive_that_expands_enormously_is_refused_without_unpacking_it() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * (import_file.MAX_UNPACKED_BYTES + 1))
    assert len(buffer.getvalue()) < import_file.MAX_FILE_BYTES  # small to send, huge to unpack
    assert "too large once unpacked" in refused(buffer.getvalue())


def test_an_archive_with_very_many_parts_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for number in range(import_file.MAX_ARCHIVE_MEMBERS + 1):
            archive.writestr(f"part{number}.xml", "")
    assert "too large once unpacked" in refused(buffer.getvalue())


def test_a_file_over_the_size_cap_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(import_file, "MAX_FILE_BYTES", 100)
    assert "larger than" in refused(b"kind\n" + b"CLASS\n" * 50)


def test_too_many_rows_are_refused_in_both_formats(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(import_file, "MAX_ROWS", 3)
    rows = [["kind"]] + [["CLASS"]] * 4
    assert "more than 3 rows" in refused(xlsx(*rows))
    assert "more than 3 rows" in refused(b"kind\n" + b"CLASS\n" * 4)
    read_table(xlsx(*rows[:4]))  # exactly three rows is fine


def test_a_spreadsheet_that_declares_a_huge_sheet_of_blank_rows_is_refused(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(import_file, "MAX_ROWS", 10)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "kind"
    sheet["A500"] = "CLASS"  # nothing between: blank rows still count
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert "more than 10 rows" in refused(buffer.getvalue())


def test_a_very_long_cell_is_refused_with_its_row_number() -> None:
    assert "Row 2" in refused(b"kind,path\nCLASS," + b"x" * 2001 + b"\n")


@pytest.mark.parametrize(
    ("data", "reason"),
    [
        (b"\n\n", "no rows|first row"),
        (b",,\nCLASS\n", "first row"),
        (b"kind,Kind\nCLASS,CLASS\n", "same name"),
    ],
)
def test_a_file_without_a_usable_header_is_refused(data: bytes, reason: str) -> None:
    import re

    assert re.search(reason, refused(data))


def test_a_row_wider_than_the_header_is_refused_in_both_formats() -> None:
    # An unquoted comma inside a value would otherwise drop the end of the row without a word.
    assert "Row 2" in refused(b"kind,path\nCLASS,Class 11,extra\n")
    assert "Row 2" in refused(xlsx(["kind", "path"], ["CLASS", "Class 11", "extra"]))
    # Empty cells beyond the header are harmless.
    assert read_table(b"kind,path\nCLASS,Class 11,,\n").rows[0].cells == {
        "kind": "CLASS",
        "path": "Class 11",
    }


def test_content_in_columns_beyond_the_limit_is_refused_even_for_a_formula() -> None:
    blanks = [None] * (import_file.MAX_COLUMNS - 1)
    for position in (0, 1, 20):  # the 41st column and some way past it
        assert "more than 40 columns" in refused(
            xlsx(["kind"], ["CLASS", *blanks, *[None] * position, "x"])
        )
        assert "formulas" in refused(xlsx(["kind"], ["CLASS", *blanks, *[None] * position, "=1+1"]))
    read_table(xlsx(["kind"], ["CLASS", *blanks]))  # empty cells out there do not matter


def test_a_spreadsheet_whose_sheet_xml_is_damaged_is_refused_not_a_server_error() -> None:
    good = io.BytesIO(xlsx(["kind"], ["CLASS"]))
    damaged = io.BytesIO()
    with zipfile.ZipFile(good) as source, zipfile.ZipFile(damaged, "w") as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename == "xl/worksheets/sheet1.xml":
                content = b"<worksheet><sheetData><row><c r="  # cut off in the middle
            target.writestr(member, content)
    assert "valid .xlsx" in refused(damaged.getvalue())


def test_a_spreadsheet_cannot_hide_cells_by_declaring_a_small_size() -> None:
    good = io.BytesIO(xlsx(["kind"], ["CLASS", *[None] * 59, "hidden"]))
    lying = io.BytesIO()
    with zipfile.ZipFile(good) as source, zipfile.ZipFile(lying, "w") as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb"<dimension[^>]*/>", b'<dimension ref="A1:B2"/>', content)
            target.writestr(member, content)
    assert "more than 40 columns" in refused(lying.getvalue())


@pytest.mark.parametrize("data", [b"kind," + b"x" * 2001 + b"\nCLASS,a\n"])
def test_the_header_cells_have_the_length_cap_too(data: bytes) -> None:
    assert "column name" in refused(data)
    assert "column name" in refused(xlsx(["kind", "x" * 2001], ["CLASS", "a"]))


def test_empty_cells_far_to_the_right_do_not_make_the_reader_walk_them() -> None:
    """A style-only cell in the last column (XFD) of every row used to make the reader build all
    16,384 cells of every row: a thousand rows took seconds, 20,000 rows would take minutes."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["kind", "path"])
    for number in range(2, 1502):
        sheet.append(["CLASS", f"Class {number}"])
        sheet.cell(row=number, column=16384).number_format = "0.00"  # empty, but styled
    buffer = io.BytesIO()
    workbook.save(buffer)
    started = time.monotonic()
    table = read_table(buffer.getvalue())
    assert time.monotonic() - started < 1.0
    assert len(table.rows) == 1500


def test_a_value_far_to_the_right_is_still_refused_among_styled_empty_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["kind"])
    sheet.append(["CLASS"])
    sheet.cell(row=2, column=16384).number_format = "0.00"
    sheet.cell(row=2, column=5000).value = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert "more than 40 columns" in refused(buffer.getvalue())


def test_a_value_in_a_column_without_a_header_name_is_refused() -> None:
    assert "no name" in refused(b"kind,,path\nCLASS,unexpected,Class 11\n")
    assert "no name" in refused(xlsx(["kind", None, "path"], ["CLASS", "unexpected", "Class 11"]))
    table = read_table(b"kind,,path\nCLASS,,Class 11\n")  # an empty unnamed column is harmless
    assert table.rows[0].cells == {"kind": "CLASS", "path": "Class 11"}


def with_sheet_xml(data: bytes, sheet_xml: bytes) -> bytes:
    """The same workbook with its first sheet's XML replaced."""
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(result, "w") as target:
        for member in source.infolist():
            content = (
                sheet_xml if member.filename == "xl/worksheets/sheet1.xml" else source.read(member)
            )
            target.writestr(member, content)
    return result.getvalue()


def packed_row_workbook(cells: int) -> bytes:
    sheet = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        b'<row r="1"><c r="A1" t="inlineStr"><is><t>kind</t></is></c></row>'
        b'<row r="2"><c r="A2" t="inlineStr"><is><t>CLASS</t></is></c>'
        + b'<c r="B2" s="1"/>' * cells
        + b"</row></sheetData></worksheet>"
    )
    return with_sheet_xml(xlsx(["kind"], ["CLASS"]), sheet)


def test_a_row_packed_with_empty_cells_is_refused_without_using_much_memory() -> None:
    """Such a row fits in a small compressed file but would become hundreds of megabytes of
    objects. A row cannot hold more cells than a sheet has columns."""
    data = packed_row_workbook(400_000)
    assert len(data) < import_file.MAX_FILE_BYTES
    tracemalloc.start()
    try:
        message = refused(data)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert "more cells than a sheet can hold" in message
    # The spreadsheet library alone peaks near 33 MB on this 7 MB sheet; keeping the cells in
    # memory as well (the earlier behaviour) took over 200 MB.
    assert peak < 64 * 1024 * 1024


def test_a_row_may_hold_as_many_cells_as_the_limit_and_not_one_more(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(import_file, "MAX_CELLS_PER_ROW", 6)
    assert [r.cells for r in read_table(packed_row_workbook(5)).rows] == [{"kind": "CLASS"}]
    assert "more cells than a sheet can hold" in refused(packed_row_workbook(6))


@pytest.mark.parametrize(
    "data",
    [
        b'kind,path\nCLASS,"Class 11"oops\n',  # a closing quote followed by more text
        b'kind,path\nCLASS,"Class 11\nSUBJECT,Physics\n',  # a quote that is never closed
    ],
)
def test_a_csv_with_broken_quoting_is_refused_not_repaired(data: bytes) -> None:
    assert "quote" in refused(data)


def sheet_with_rows(rows: bytes) -> bytes:
    sheet = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>kind</t></is></c></row>'
        + rows
        + b"</sheetData></worksheet>"
    )
    return with_sheet_xml(xlsx(["kind"], ["CLASS"]), sheet)


def test_millions_of_empty_rows_are_refused_early_without_using_much_memory() -> None:
    data = sheet_with_rows(b"<row/>" * 2_500_000)  # about 15 MB unpacked, tiny once compressed
    assert len(data) < import_file.MAX_FILE_BYTES
    tracemalloc.start()
    try:
        message = refused(data)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert "more than 20,000 rows" in message
    assert peak < 64 * 1024 * 1024


def test_a_row_far_down_the_sheet_is_refused_without_walking_the_gap() -> None:
    assert "more than 20,000 rows" in refused(sheet_with_rows(b'<row r="999999"/>'))


def moved_sheet(data: bytes, sheet_xml: bytes | None = None) -> bytes:
    """The first sheet stored at xl/custom.xml, with the workbook's relationship pointing there."""
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(result, "w") as target:
        for member in source.infolist():
            content = source.read(member)
            name = member.filename
            if name == "xl/worksheets/sheet1.xml":
                name, content = "xl/custom.xml", sheet_xml or content
            elif name == "xl/_rels/workbook.xml.rels":
                content = content.replace(b"worksheets/sheet1.xml", b"custom.xml")
                content = content.replace(b"/xl/custom.xml", b"custom.xml")
            target.writestr(name, content)
    return result.getvalue()


def test_the_sheet_is_found_through_the_workbook_relationships_wherever_it_is_stored() -> None:
    good = xlsx(["kind", "path"], ["CLASS", "Class 11"])
    assert [r.cells for r in read_table(moved_sheet(good)).rows] == [
        {"kind": "CLASS", "path": "Class 11"}
    ]
    far_cell = xlsx(["kind"], ["CLASS", *[None] * 59, "hidden"])
    assert "more than 40 columns" in refused(moved_sheet(far_cell))


def test_a_sheet_with_hundreds_of_thousands_of_other_elements_does_not_pile_up_in_memory() -> None:
    merges = b'<mergeCell ref="A1:B2"/>' * 400_000  # about 10 MB of XML, small once compressed
    sheet = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        b'<row r="1"><c r="A1" t="inlineStr"><is><t>kind</t></is></c></row>'
        b'<row r="2"><c r="A2" t="inlineStr"><is><t>CLASS</t></is></c></row></sheetData>'
        b"<mergeCells>" + merges + b"</mergeCells></worksheet>"
    )
    data = with_sheet_xml(xlsx(["kind"], ["CLASS"]), sheet)
    tracemalloc.start()
    try:
        table = read_table(data)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert [r.cells for r in table.rows] == [{"kind": "CLASS"}]
    assert peak < 20 * 1024 * 1024


def test_rows_out_of_order_or_repeated_are_refused() -> None:
    row = b'<row r="%d"><c r="A%d" t="inlineStr"><is><t>x</t></is></c></row>'
    assert "not in order" in refused(sheet_with_rows(row % (3, 3) + row % (2, 2)))
    assert "not in order" in refused(sheet_with_rows(b'<row r="1"/>'))  # the header is row 1


def test_values_of_every_kind_of_cell_are_read_as_text() -> None:
    sheet = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        b'<row r="1"><c r="A1" t="inlineStr"><is><t>kind</t></is></c>'
        b'<c r="B1" t="inlineStr"><is><t>n</t></is></c><c r="C1" t="inlineStr"><is><t>b</t></is>'
        b'</c><c r="D1" t="inlineStr"><is><t>s</t></is></c></row>'
        b'<row r="2"><c r="A2" t="inlineStr"><is><r><t>CL</t></r><r><t>ASS</t></r>'
        b"<rPh><t>ignored</t></rPh></is></c>"
        b'<c r="B2"><v>3.0</v></c><c r="C2" t="b"><v>1</v></c><c r="D2" t="s"><v>0</v></c></row>'
        b"</sheetData></worksheet>"
    )
    strings = (
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<si><t> shared </t></si></sst>"
    )
    result = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(xlsx(["kind"], ["CLASS"]))) as source,
        zipfile.ZipFile(result, "w") as target,
    ):
        for member in source.infolist():
            content = source.read(member)
            if member.filename == "xl/worksheets/sheet1.xml":
                content = sheet
            if member.filename != "xl/sharedStrings.xml":
                target.writestr(member, content)
        target.writestr("xl/sharedStrings.xml", strings)
    table = read_table(result.getvalue())
    assert table.rows[0].cells == {"kind": "CLASS", "n": "3", "b": "TRUE", "s": "shared"}


def test_one_cell_with_hundreds_of_thousands_of_children_is_refused_without_using_much_memory() -> (
    None
):
    cell = b'<c r="A2" t="inlineStr"><is>' + b"<r/>" * 400_000 + b"</is></c>"
    data = sheet_with_rows(b'<row r="2">' + cell + b"</row>")
    tracemalloc.start()
    try:
        message = refused(data)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert "too complex" in message
    assert peak < 20 * 1024 * 1024


def test_a_shared_string_with_hundreds_of_thousands_of_parts_is_refused() -> None:
    strings = b"<sst><si>" + b"<r/>" * 400_000 + b"</si></sst>"
    assert "too complex" in refused(
        with_part(xlsx(["kind"], ["CLASS"]), "xl/sharedStrings.xml", strings)
    )


def test_a_shared_string_index_outside_the_table_is_refused() -> None:
    strings = b"<sst><si><t>only</t></si></sst>"
    for index in (b"-1", b"1", b"x"):
        row = b'<row r="2"><c r="A2" t="s"><v>' + index + b"</v></c></row>"
        data = with_part(sheet_with_rows(row), "xl/sharedStrings.xml", strings)
        assert "valid .xlsx" in refused(data)


def with_part(data: bytes, name: str, content: bytes) -> bytes:
    """The workbook with one part added or replaced."""
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(result, "w") as target:
        for member in source.infolist():
            if member.filename != name:
                target.writestr(member, source.read(member))
        target.writestr(name, content)
    return result.getvalue()


def test_deeply_nested_xml_is_refused_quickly() -> None:
    depth = 200_000
    data = sheet_with_rows(b"<a>" * depth + b"</a>" * depth)  # not inside any cell
    started = time.monotonic()
    assert "too complex" in refused(data)
    assert time.monotonic() - started < 2.0


# --- Attacks named in the import threat model (docs/architecture/threat-model-import.md, K-10) ---

_BILLION_LAUGHS = (
    b'<?xml version="1.0"?><!DOCTYPE lolz ['
    b'<!ENTITY a "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa">'
    b'<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;"><!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">'
    b'<!ENTITY d "&c;&c;&c;&c;&c;&c;&c;&c;&c;&c;"><!ENTITY e "&d;&d;&d;&d;&d;&d;&d;&d;&d;&d;">'
    b'<!ENTITY f "&e;&e;&e;&e;&e;&e;&e;&e;&e;&e;"><!ENTITY g "&f;&f;&f;&f;&f;&f;&f;&f;&f;&f;">'
    b']><worksheet><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&g;</t></is></c>'
    b"</row></sheetData></worksheet>"
)
_FILE_ENTITY = (
    b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]>'
    b'<worksheet><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&secret;</t></is></c>'
    b"</row></sheetData></worksheet>"
)
_URL_ENTITY = _FILE_ENTITY.replace(b"file:///etc/passwd", b"http://127.0.0.1:9/leak")


@pytest.mark.parametrize("sheet", [_BILLION_LAUGHS, _FILE_ENTITY, _URL_ENTITY])
def test_xml_entities_are_refused_in_a_sheet_and_nothing_is_read_or_fetched(sheet: bytes) -> None:
    """The "billion laughs" expansion and external entities (a local file, a web address) are the
    classic attacks on XML readers. defusedxml refuses any entity declaration, so none of them is
    expanded, read or fetched."""
    started = time.monotonic()
    message = refused(with_part(xlsx(["kind"], ["CLASS"]), "xl/worksheets/sheet1.xml", sheet))
    assert "valid .xlsx" in message and "root:" not in message
    assert time.monotonic() - started < 2.0


def test_xml_entities_are_refused_in_the_shared_strings_too() -> None:
    data = with_part(xlsx(["kind"], ["CLASS"]), "xl/sharedStrings.xml", _BILLION_LAUGHS)
    assert "valid .xlsx" in refused(data)


@pytest.mark.parametrize(
    "target",
    [
        "../../../etc/passwd",  # up and out of the archive
        "/etc/passwd",  # an absolute path
        "http://127.0.0.1:9/sheet1.xml",  # a web address
        "file:///etc/passwd",
    ],
)
def test_a_sheet_relationship_that_points_outside_the_archive_is_refused(target: str) -> None:
    """The reader only ever looks a part up by name inside the archive in memory. It never opens a
    file path or a web address taken from the workbook."""
    good = xlsx(["kind"], ["CLASS"])
    with zipfile.ZipFile(io.BytesIO(good)) as archive:
        rels = archive.read("xl/_rels/workbook.xml.rels").decode()
    pointed = rels.replace('Target="/xl/worksheets/sheet1.xml"', f'Target="{target}"')
    assert pointed != rels
    assert "valid .xlsx" in refused(with_part(good, "xl/_rels/workbook.xml.rels", pointed.encode()))


def test_an_archive_that_understates_the_size_of_a_part_cannot_slip_more_data_past_the_cap() -> (
    None
):
    """The unpacked-size check trusts the sizes written in the archive, which an attacker writes.
    The zip reader stops at the declared size and then fails its own checks, so a lie gives a
    refusal, not extra data."""
    honest = bytearray(
        sheet_with_rows(b'<row r="2"><c r="A2" t="inlineStr"><is><t>x</t></is></c></row>')
    )
    assert read_table(bytes(honest)).rows  # the same file is fine until the sizes are falsified
    lying = bytearray(honest)
    for match in re.finditer(rb"PK\x01\x02", bytes(lying)):  # each part's entry in the directory
        position = match.start() + 24  # its "uncompressed size" field
        lying[position : position + 4] = (10).to_bytes(4, "little")
    assert "valid .xlsx" in refused(bytes(lying))


def test_a_csv_of_millions_of_commas_is_refused_before_it_is_split_into_fields() -> None:
    data = b"," * 5_000_000
    tracemalloc.start()
    try:
        message = refused(data)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert "too many cells" in message
    assert peak < 20 * 1024 * 1024


def test_an_element_with_millions_of_attributes_is_refused_before_it_is_built() -> None:
    attributes = b" ".join(b'a%d="1"' % number for number in range(1_500_000))  # ~15 MB, no parser
    data = sheet_with_rows(b"<odd " + attributes + b"/>")
    tracemalloc.start()
    try:
        message = refused(data)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert "too complex" in message
    assert peak < 20 * 1024 * 1024


def test_a_tag_split_across_reads_is_still_measured() -> None:
    long_tag = b"<odd " + b" ".join(b'a%d="1"' % number for number in range(1_000)) + b"/>"
    guard = import_file._TagGuard(io.BytesIO(b"<a>" + b"x" * 3_000 + long_tag))
    with pytest.raises(ImportFileInvalid, match="too complex"):
        while guard.read(500):  # the tag is cut into many reads, none holding all of it
            pass
    short = import_file._TagGuard(io.BytesIO(b"<a " + b'b="1" ' * 100 + b"/>" * 1))
    assert short.read(100) + short.read(100000)  # an ordinary tag passes, read in pieces
